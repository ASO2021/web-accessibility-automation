"""Tests sin navegador: validan que el resumen y los reportes se generen bien."""
import json

import pytest
from openpyxl import load_workbook

from config import settings
from utils.report import (build_html, format_scan_results, safe_name, save_reports,
                          build_contrast_css, contrast_recommendations,
                          contrast_color_audit, _contrast_ratio, _hex_rgb, serious_or_critical,
                          summarize, wcag_refs)

SAMPLE = {
    "url": "https://example.com/",
    "testEngine": {"version": "4.14.0"},
    "passes": [{"id": "html-has-lang"}, {"id": "document-title"}],
    "violations": [
        {
            "id": "image-alt", "impact": "critical", "help": "Images must have alternate text",
            "helpUrl": "https://dequeuniversity.com/rules/axe/4.14/image-alt",
            "tags": ["cat.text-alternatives", "wcag2a", "wcag111"],
            "nodes": [{"target": ["img.logo"]}, {"target": ["img.banner"]}],
        },
        {
            "id": "color-contrast", "impact": "serious", "help": "Elements must have enough contrast",
            "helpUrl": "https://dequeuniversity.com/rules/axe/4.14/color-contrast",
            "tags": ["wcag2aa", "wcag143"],
            "nodes": [{"target": ["p.note"]}],
        },
        {
            "id": "region", "impact": "moderate", "help": "Content should be in landmarks",
            "helpUrl": "https://dequeuniversity.com/rules/axe/4.14/region",
            "tags": ["best-practice"], "nodes": [{"target": ["div"]}],
        },
    ],
}


@pytest.mark.unit
def test_summarize_counts_by_impact():
    summary = summarize(SAMPLE)

    assert summary["rules_violated"] == 3
    assert summary["elements_affected"] == 4
    assert summary["rules_passed"] == 2
    assert summary["by_impact"] == {"critical": 1, "serious": 1, "moderate": 1, "minor": 0}


@pytest.mark.unit
def test_format_scan_results_shows_violation_details():
    output = format_scan_results(SAMPLE)

    assert "[CRITICAL] image-alt" in output
    assert "WCAG: wcag2a, wcag111" in output
    assert "Elemento 1: img.logo" in output


@pytest.mark.unit
def test_format_scan_results_includes_contrast_failure_details():
    contrast_violation = {
        "url": "https://example.com/",
        "violations": [{
            "id": "color-contrast",
            "impact": "serious",
            "help": "Elements must have sufficient color contrast",
            "tags": ["wcag2aa", "wcag143"],
            "nodes": [{
                "target": ["p.notice"],
                "any": [{
                    "id": "color-contrast",
                    "data": {
                        "fgColor": "#41545d",
                        "bgColor": "#a9b8bf",
                        "contrastRatio": 3.88,
                        "expectedContrastRatio": "4.5:1",
                    },
                }],
                "failureSummary": "Fix any of the following:\nElement has insufficient color contrast.",
            }],
        }],
    }

    output = format_scan_results(contrast_violation)

    assert "[SERIOUS] color-contrast" in output
    assert "Elemento 1: p.notice" in output
    assert "insufficient color contrast" in output
    assert "Color de texto sugerido:" in output
    assert "4.5:1" in output


@pytest.mark.unit
def test_contrast_recommendation_meets_wcag_ratio_and_generates_css():
    contrast_violation = {
        "violations": [{
            "id": "color-contrast",
            "nodes": [{
                "target": ["p.notice"],
                "any": [{
                    "id": "color-contrast",
                    "data": {
                        "fgColor": "#41545d",
                        "bgColor": "#a9b8bf",
                        "contrastRatio": 3.88,
                        "expectedContrastRatio": "4.5:1",
                    },
                }],
            }],
        }],
    }

    recommendations = contrast_recommendations(contrast_violation)
    css = build_contrast_css(contrast_violation)

    assert len(recommendations) == 1
    assert recommendations[0]["status"] == "css_available"
    assert recommendations[0]["suggested_foreground"]
    assert _contrast_ratio(
        _hex_rgb(recommendations[0]["suggested_foreground"]),
        _hex_rgb(recommendations[0]["background"]),
    ) >= recommendations[0]["minimum_ratio"]
    assert "p.notice {" in css
    assert f"color: {recommendations[0]['suggested_foreground']} !important;" in css


@pytest.mark.unit
def test_contrast_color_audit_includes_pass_fail_and_manual_review():
    results = {
        "passes": [{
            "id": "color-contrast",
            "nodes": [{
                "target": ["p.good"],
                "any": [{
                    "id": "color-contrast",
                    "data": {
                        "fgColor": "#000000",
                        "bgColor": "#ffffff",
                        "contrastRatio": 21,
                        "expectedContrastRatio": "4.5:1",
                    },
                }],
            }],
        }],
        "violations": [{
            "id": "color-contrast",
            "nodes": [{
                "target": ["p.bad"],
                "any": [{
                    "id": "color-contrast",
                    "data": {
                        "fgColor": "#777777",
                        "bgColor": "#888888",
                        "contrastRatio": 1.26,
                        "expectedContrastRatio": "4.5:1",
                    },
                }],
            }],
        }],
        "incomplete": [{
            "id": "color-contrast",
            "nodes": [{"target": ["p.unknown"], "any": []}],
        }],
    }

    audit = contrast_color_audit(results)

    assert [(item["selector"], item["status"]) for item in audit] == [
        ("p.good", "conforme"),
        ("p.bad", "error"),
        ("p.unknown", "revision_manual"),
    ]


@pytest.mark.unit
def test_serious_or_critical_filters_rules():
    ids = [v["id"] for v in serious_or_critical(SAMPLE)]

    assert ids == ["image-alt", "color-contrast"]


@pytest.mark.unit
def test_wcag_refs_keeps_only_wcag_tags():
    assert wcag_refs(SAMPLE["violations"][0]) == ["wcag111", "wcag2a"]


@pytest.mark.unit
def test_safe_name_removes_special_characters():
    assert safe_name("https://www.saucedemo.com/inventory.html") == "https_www_saucedemo_com_inventory_html"
    assert safe_name("///") == "report"


@pytest.mark.unit
def test_html_report_lists_rules_and_escapes_content():
    data = {**SAMPLE, "violations": [{
        **SAMPLE["violations"][0], "help": "<script>alert(1)</script>"}]}

    page = build_html(data, "Home")

    assert "image-alt" in page
    assert "<script>alert(1)</script>" not in page
    assert "&lt;script&gt;" in page


@pytest.mark.unit
def test_html_report_with_no_violations():
    page = build_html({"url": "x", "violations": [], "passes": []}, "Clean page")

    assert "No violations found." in page


@pytest.mark.unit
def test_save_reports_writes_json_and_html(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "REPORTS_DIR", str(tmp_path))

    json_path, html_path = save_reports(SAMPLE, "My Page", "My Page")

    assert json_path.exists() and html_path.exists()
    assert json_path.with_suffix(".suggestions.css").exists()
    excel_path = json_path.with_suffix(".xlsx")
    assert excel_path.exists()
    assert excel_path.name in html_path.read_text(encoding="utf-8")
    saved = json.loads(json_path.read_text(encoding="utf-8"))
    assert saved["summary"]["rules_violated"] == 3
    workbook = load_workbook(excel_path, read_only=True)
    assert workbook.sheetnames == ["Resumen", "Hallazgos", "Contraste", "Paleta"]
    assert workbook["Hallazgos"].max_row == 5
    assert workbook["Hallazgos"]["A1"].value == "Impacto"
    assert workbook["Contraste"]["A1"].value == "Elemento afectado (selector)"
    assert workbook["Paleta"]["A1"].value == "Resultado"
    workbook.close()
